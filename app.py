"""API local: recebe imagem, devolve o que foi lido. Sem login, sem banco, sem nuvem."""
from __future__ import annotations

import csv
import io
import mimetypes
import threading
import webbrowser
from contextlib import asynccontextmanager
from pathlib import Path

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel

import ocr_engine

BASE = Path(__file__).parent
MAX_BYTES = 15 * 1024 * 1024

# No Windows o Python lê o MIME do registro, e em várias máquinas ".js" vem como
# "text/plain". O navegador recusa <script type="module"> com esse MIME e o front não carrega.
mimetypes.add_type("text/javascript", ".js")
mimetypes.add_type("text/css", ".css")
mimetypes.add_type("font/woff2", ".woff2")

@asynccontextmanager
async def lifespan(_: FastAPI):
    # Carrega os modelos uma vez, para a primeira leitura do usuário não pagar esse custo.
    ocr_engine.get_engine()
    yield


app = FastAPI(title="OCR Buddies", docs_url=None, redoc_url=None, lifespan=lifespan)


class TablePayload(BaseModel):
    rows: list[list[str]]


@app.post("/api/ocr")
async def ocr(file: UploadFile = File(...)):
    data = await file.read()
    if len(data) > MAX_BYTES:
        raise HTTPException(413, "Imagem maior que 15 MB.")
    try:
        return ocr_engine.analyze(data)
    except Exception as exc:  # imagem corrompida / formato não suportado
        raise HTTPException(400, f"Não consegui abrir essa imagem: {exc}") from exc


@app.post("/api/export/csv")
def export_csv(payload: TablePayload):
    buf = io.StringIO()
    csv.writer(buf, delimiter=";").writerows(payload.rows)  # ';' abre direto no Excel pt-BR
    body = "﻿" + buf.getvalue()  # BOM: acentos corretos no Excel
    return Response(
        body.encode("utf-8"),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="tabela.csv"'},
    )


def _to_number(cell: str):
    """'1.234,56' -> 1234.56 ; mantém texto quando não é número."""
    s = cell.strip().replace("R$", "").replace("%", "").replace(" ", "")
    if not s or not any(ch.isdigit() for ch in s):
        return cell
    try:
        if "," in s:
            return float(s.replace(".", "").replace(",", "."))
        if s.count(".") > 1:
            return float(s.replace(".", ""))
        return float(s) if "." in s else int(s)
    except ValueError:
        return cell


@app.post("/api/export/xlsx")
def export_xlsx(payload: TablePayload):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tabela"
    for r, row in enumerate(payload.rows, start=1):
        for c, cell in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=_to_number(cell))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), 60)
    out = io.BytesIO()
    wb.save(out)
    return Response(
        out.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="tabela.xlsx"'},
    )


@app.get("/")
def index():
    return FileResponse(BASE / "static" / "index.html")


app.mount("/static", StaticFiles(directory=BASE / "static"), name="static")

if __name__ == "__main__":
    import uvicorn

    URL = "http://127.0.0.1:8000"
    threading.Timer(2.5, lambda: webbrowser.open(URL)).start()  # abre o navegador sozinho
    uvicorn.run(app, host="127.0.0.1", port=8000)  # só local: nada exposto na rede
